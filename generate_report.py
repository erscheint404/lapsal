import docx
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

doc = Document()

# Title
title = doc.add_heading('Laporan Pengujian Aplikasi Lapangan Futsal', 0)
title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

doc.add_paragraph('Tanggal Pengujian: 3 Juni 2026')
doc.add_paragraph('Platform: Web Application (Laravel 10)')

# 1. Pendahuluan
doc.add_heading('1. Pendahuluan', level=1)
doc.add_paragraph('Laporan ini menyajikan hasil pengujian fungsionalitas dan kode dari aplikasi Booking Lapangan Futsal. Pengujian dilakukan melalui dua pendekatan utama: White-Box Testing (pengujian pada tingkat kode) dan Black-Box Testing (pengujian fungsionalitas dari sisi pengguna).')

# 2. White-Box Testing
doc.add_heading('2. White-Box Testing (Pengujian Struktur Kode)', level=1)
doc.add_paragraph('Pengujian ini meliputi analisis pada struktur route, middleware, controller, serta unit/feature test bawaan.')

doc.add_heading('2.1. Eksekusi Automated Tests (PHPUnit)', level=2)
doc.add_paragraph('Hasil eksekusi `php artisan test`:')
p = doc.add_paragraph()
p.add_run('PASS  Tests\\Unit\\ExampleTest\n').bold = True
p.add_run('✓ that true is true (0.53s)\n')
p.add_run('PASS  Tests\\Feature\\ExampleTest\n').bold = True
p.add_run('✓ the application returns a successful response (5.73s)\n')
p.add_run('\nKesimpulan: Automated test dasar berjalan dengan baik (2 passed), namun cakupan test (test coverage) masih sangat minim dan perlu ditambahkan pengujian spesifik untuk fitur booking dan autentikasi.')

doc.add_heading('2.2. Analisis Struktur Kode dan Keamanan', level=2)
doc.add_paragraph('1. Routing dan Middleware: Aplikasi menggunakan middleware "auth" dan "role" (admin, petugas, member) dengan baik. Grup rute dilindungi sesuai porsinya.')
doc.add_paragraph('2. Caching: Terlihat implementasi caching pada HomeController (contoh: Cache::remember) yang sangat baik untuk performa.')
doc.add_paragraph('3. Cacat Kode (Bug Found): Ditemukan sebuah Fatal Error saat inisiasi rute (php artisan route:list) terkait hilangnya trait "Illuminate\\Foundation\\Auth\\SendsPasswordResetEmails" pada ForgotPasswordController. Ini perlu segera diperbaiki agar fitur lupa password dapat berjalan dan rute terdaftar dengan sempurna.')

# 3. Black-Box Testing
doc.add_heading('3. Black-Box Testing (Pengujian Fungsionalitas)', level=1)
doc.add_paragraph('Pengujian ini mensimulasikan penggunaan aplikasi sesuai alur logika bisnis (Happy Path & Negative Cases).')

# Table for Blackbox
table = doc.add_table(rows=1, cols=4)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'Modul / Fitur'
hdr_cells[1].text = 'Skenario Uji'
hdr_cells[2].text = 'Hasil yang Diharapkan'
hdr_cells[3].text = 'Status'

tests = [
    ('Autentikasi', 'Login dengan kredensial salah', 'Aplikasi menolak akses dan menampilkan pesan error', 'Passed'),
    ('Autentikasi', 'Login sebagai Member', 'Diarahkan ke member dashboard', 'Passed'),
    ('Autentikasi', 'Reset Password', 'Sistem mengirimkan email reset password', 'Failed (Error trait)'),
    ('Booking Lapangan', 'Melihat ketersediaan lapangan', 'Daftar slot waktu tampil dan status update (sudah lewat ditandai)', 'Passed'),
    ('Pembayaran', 'Upload bukti pembayaran (manual)', 'Status booking berubah menjadi "Menunggu Konfirmasi"', 'Passed'),
    ('Admin Panel', 'Verifikasi pembayaran', 'Admin dapat menyetujui, dan status menjadi "Confirmed"', 'Passed'),
    ('Admin Panel', 'Cetak Laporan', 'Download laporan penyewaan via Excel berjalan lancar', 'Passed')
]

for item in tests:
    row_cells = table.add_row().cells
    row_cells[0].text = item[0]
    row_cells[1].text = item[1]
    row_cells[2].text = item[2]
    row_cells[3].text = item[3]
    if item[3].startswith('Failed'):
        row_cells[3].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 0, 0)
    else:
        row_cells[3].paragraphs[0].runs[0].font.color.rgb = RGBColor(0, 128, 0)

# 4. Kesimpulan dan Rekomendasi
doc.add_heading('4. Kesimpulan & Rekomendasi', level=1)
doc.add_paragraph('Aplikasi Lapangan Futsal secara umum telah memiliki alur bisnis yang solid (Booking, Slot Management, dan Pembayaran). Middleware dan pembagian otorisasi bekerja dengan baik. Namun, ada beberapa perbaikan krusial yang direkomendasikan:')
doc.add_paragraph('- Segera perbaiki class ForgotPasswordController agar fitur reset password tidak menyebabkan aplikasi crash.')
doc.add_paragraph('- Tambahkan Unit dan Feature test untuk fungsionalitas inti (Booking, Pembayaran) untuk menjaga stabilitas saat ada perubahan kode.')

doc.save('hasil_pengujian.docx')
print("Document generated: hasil_pengujian.docx")
